#!/usr/bin/env python3
"""Сутрешен вестник: събира RSS новини, обобщава ги с Claude и прави Word (.docx) вестник."""

import calendar
import difflib
import html
import json
import logging
import re
import shutil
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from string import Template

import feedparser
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent
NEWSPAPERS = BASE / "newspapers"
LOGS = BASE / "logs"
SECTIONS = ["България", "Армения", "Свят", "Технологии и AI", "Любопитно"]
USER_AGENT = "Mozilla/5.0 (Macintosh) SutreshenVestnik/1.0"

FONT = "PT Serif"   # шрифт, създаден за кирилица; има го на всеки Mac
INK = RGBColor(0x1C, 0x1A, 0x17)
ACCENT = RGBColor(0x8B, 0x1E, 0x1E)
MUTED = RGBColor(0x6B, 0x64, 0x5A)

WEEKDAYS = ["понеделник", "вторник", "сряда", "четвъртък", "петък", "събота", "неделя"]
MONTHS = ["януари", "февруари", "март", "април", "май", "юни", "юли",
          "август", "септември", "октомври", "ноември", "декември"]

log = logging.getLogger("vestnik")


def setup_logging():
    LOGS.mkdir(exist_ok=True)
    fmt = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
    file_handler = logging.FileHandler(LOGS / "vestnik.log", encoding="utf-8")
    file_handler.setFormatter(fmt)
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(fmt)
    log.addHandler(file_handler)
    log.addHandler(console)
    log.setLevel(logging.INFO)


def load_config():
    with open(BASE / "sources.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f)


# ---------- 1. Събиране ----------

def clean_text(text, limit=300):
    text = re.sub(r"<[^>]+>", " ", text or "")
    text = html.unescape(re.sub(r"\s+", " ", text)).strip()
    return text[:limit] + ("…" if len(text) > limit else "")


def fetch_source(source, since, limit):
    """Връща списък с новини от един източник. При грешка връща празен списък."""
    try:
        req = urllib.request.Request(source["url"], headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=20) as resp:
            feed = feedparser.parse(resp.read())
        items = []
        for entry in feed.entries:
            parsed = entry.get("published_parsed") or entry.get("updated_parsed")
            if parsed and calendar.timegm(parsed) < since:
                continue
            title = clean_text(entry.get("title"), 200)
            if not title or not entry.get("link"):
                continue
            items.append({
                "title": title,
                "summary": clean_text(entry.get("summary")),
                "link": entry.get("link"),
                "source": source["name"],
                "section": source.get("section", ""),
            })
            if len(items) >= limit:
                break
        log.info("OK   %-14s %d новини", source["name"], len(items))
        return items
    except Exception as e:
        log.error("ГРЕШКА %-14s %s: %s", source["name"], source["url"], e)
        return []


def fetch_all(config):
    settings = config["settings"]
    since = time.time() - settings["hours_back"] * 3600
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = pool.map(
            lambda s: fetch_source(s, since, settings["per_source_limit"]),
            config["sources"])
    return [item for items in results for item in items]


# ---------- 2. Махане на очевидни дубликати ----------

def normalize(title):
    return re.sub(r"[^\w ]", "", title.lower())


def dedupe(items):
    kept = []
    for item in items:
        norm = normalize(item["title"])
        if any(difflib.SequenceMatcher(None, norm, normalize(k["title"])).ratio() > 0.85
               for k in kept):
            continue
        kept.append(item)
    return kept


# ---------- 3. Обобщаване с Claude ----------

PROMPT = """Ти си главен редактор на кратък сутрешен вестник на български език.
Читателят се интересува най-вече от: $interests.

По-долу има списък с новини от последните 24 часа. Всяка започва с номер в [квадратни скоби].

Задачи:
1. Обедини новините, които разказват една и съща история (от различни източници), в една.
2. Избери около $max_stories най-важни истории. Предпочитай значими събития пред дребни.
   Включи поне 2 за Армения и поне 2 за технологии/AI, ако има подходящи.
3. За всяка история напиши на български:
   - "title": кратко и ясно заглавие (до 12 думи);
   - "summary": неутрално и фактологично резюме, без измислици. Вестникът трябва да се
     събере на ЕДНА страница: за "top" историите 3 изречения (до 55 думи), за останалите
     2–3 изречения (до 40 думи). Всички резюмета общо — до 600 думи;
   - "section": една от рубриките: $sections;
   - "top": true за 3–5-те най-важни истории за деня, иначе false;
   - "ids": номерата на всички използвани новини от списъка.
4. Напиши "day_in_three": обобщение на деня точно в 3 кратки изречения (общо до 60 думи).

Отговори САМО с валиден JSON, без обяснения и без ```, в този формат:
{"day_in_three": "...", "stories": [{"title": "...", "summary": "...", "section": "...", "top": true, "ids": [1, 5]}]}

НОВИНИ:
$news
"""


def find_claude():
    return shutil.which("claude") or str(Path.home() / ".local/bin/claude")


def extract_json(text):
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("В отговора няма JSON")
    return json.loads(text[start:end + 1])


def summarize(items, settings):
    news = "\n".join(
        "[{}] ({}, {}) {} — {}".format(i, it["source"], it["section"], it["title"], it["summary"])
        for i, it in enumerate(items))
    prompt = Template(PROMPT).substitute(
        interests=settings["interests"],
        max_stories=settings["max_stories"],
        sections=", ".join(SECTIONS),
        news=news)

    cmd = [find_claude(), "-p", "--output-format", "json",
           "--model", settings.get("claude_model", "sonnet"),
           "--tools", "", "--no-session-persistence"]
    last_error = None
    for attempt in (1, 2):
        try:
            log.info("Claude обобщава %d новини (опит %d)...", len(items), attempt)
            proc = subprocess.run(cmd, input=prompt, capture_output=True, text=True,
                                  timeout=900, cwd=tempfile.gettempdir())
            if proc.returncode != 0:
                raise RuntimeError("claude върна код {}: {}".format(
                    proc.returncode, (proc.stderr or proc.stdout)[:500]))
            outer = json.loads(proc.stdout)
            if outer.get("is_error"):
                raise RuntimeError("claude грешка: {}".format(outer.get("result")))
            data = extract_json(outer["result"])
            if not data.get("stories"):
                raise ValueError("Няма новини в отговора")
            return data
        except Exception as e:
            last_error = e
            log.warning("Опит %d неуспешен: %s", attempt, e)
    raise RuntimeError("Claude не успя да обобщи новините: {}".format(last_error))


# ---------- 4. Времето ----------

WEATHER_CODES = [
    ((0,), "Ясно"), ((1,), "Предимно ясно"), ((2,), "Променлива облачност"),
    ((3,), "Облачно"), ((45, 48), "Мъгла"), ((51, 53, 55, 56, 57), "Ръмеж"),
    ((61,), "Слаб дъжд"), ((63,), "Дъжд"), ((65,), "Силен дъжд"), ((66, 67), "Леден дъжд"),
    ((71, 73, 75, 77), "Сняг"), ((80, 81, 82), "Превалявания"),
    ((85, 86), "Снежни превалявания"), ((95, 96, 99), "Гръмотевични бури"),
]


def weather_text(code):
    for codes, text in WEATHER_CODES:
        if code in codes:
            return text
    return "Променливо"


def fetch_weather(place):
    """Прогноза за днес от Open-Meteo (безплатно, без ключ). При грешка връща None."""
    url = ("https://api.open-meteo.com/v1/forecast?latitude={lat}&longitude={lon}"
           "&daily=weather_code,temperature_2m_max,temperature_2m_min,"
           "precipitation_probability_max,wind_speed_10m_max"
           "&hourly=temperature_2m&timezone=auto&forecast_days=1").format(**place)
    try:
        req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        with urllib.request.urlopen(req, timeout=20) as resp:
            raw = json.loads(resp.read())
        daily, hourly = raw["daily"], raw["hourly"]["temperature_2m"]
        weather = {
            "city": place["name"],
            "text": weather_text(daily["weather_code"][0]),
            "max": round(daily["temperature_2m_max"][0]),
            "min": round(daily["temperature_2m_min"][0]),
            "rain": daily["precipitation_probability_max"][0] or 0,
            "wind": round(daily["wind_speed_10m_max"][0]),
            "morning": round(hourly[8]),
            "noon": round(hourly[14]),
            "evening": round(hourly[20]),
        }
        log.info("OK   Времето в %s: %s, %d°/%d°", place["name"], weather["text"],
                 weather["max"], weather["min"])
        return weather
    except Exception as e:
        log.error("ГРЕШКА времето за %s: %s", place.get("name"), e)
        return None


def weather_tip(w):
    if w["rain"] >= 50:
        return "Вземи чадър."
    if w["min"] <= 0:
        return "Облечи се топло — възможно е заледяване."
    if w["max"] >= 30:
        return "Горещо е — пий повече вода."
    if w["max"] - w["min"] >= 10:
        return "Хладна сутрин, топъл следобед — облечи се на слоеве."
    if w["wind"] >= 40:
        return "Силен вятър — внимавай навън."
    return "Приятен ден!"


# ---------- 5. Word документ ----------

def bg_date(d):
    return "{}, {} {} {} г.".format(WEEKDAYS[d.weekday()].capitalize(), d.day,
                                    MONTHS[d.month - 1], d.year)


def set_font(run, size, bold=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), font)


def spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_border(paragraph, side, size=8, color="1C1A17", style="single", space=1):
    """Линия над или под параграф (size е в осмини от точка)."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    line = OxmlElement("w:" + side)
    line.set(qn("w:val"), style)
    line.set(qn("w:sz"), str(size))
    line.set(qn("w:space"), str(space))
    line.set(qn("w:color"), color)
    borders.append(line)


def add_shading(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def add_hyperlink(paragraph, url, text, size=7.5):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "8B1E1E")
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def set_columns(section, num, space_cm=0.6):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space_cm * 567)))


def story_links(story, items):
    """По един линк от всеки източник."""
    links, seen = [], set()
    for i in story.get("ids", []):
        if isinstance(i, int) and 0 <= i < len(items) and items[i]["source"] not in seen:
            seen.add(items[i]["source"])
            links.append((items[i]["source"], items[i]["link"]))
    return links


def add_heading(doc, text, color=INK):
    p = doc.add_paragraph()
    spacing(p, before=4, after=1.5)
    add_border(p, "bottom", size=8, color=str(color))
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text.upper()), 8.5, bold=True, color=color)


def add_story(doc, story, items, big=False):
    title = doc.add_paragraph()
    spacing(title, before=3, after=0.5, line=1.0)
    title.paragraph_format.keep_with_next = True
    set_font(title.add_run(story["title"]), 10.5 if big else 9.5, bold=True)

    body = doc.add_paragraph()
    spacing(body, after=1, line=1.0)
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_font(body.add_run(story["summary"] + " "), 8.5)
    for n, (name, url) in enumerate(story_links(story, items)):
        if n:
            set_font(body.add_run(" · "), 7.5, color=MUTED)
        add_hyperlink(body, url, name)


def add_day_box(doc, data, weather):
    """Отгоре: „Денят с 3 изречения“ вляво и времето за деня вдясно."""
    table = doc.add_table(rows=1, cols=2 if weather else 1)
    table.autofit = False
    widths = [Cm(12.2), Cm(6.8)] if weather else [Cm(19)]
    for cell, width in zip(table.rows[0].cells, widths):
        cell.width = width
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    left = table.rows[0].cells[0]
    add_shading(left, "F3EEE3")
    label = left.paragraphs[0]
    spacing(label, before=4, after=2)
    set_font(label.add_run("ДЕНЯТ С 3 ИЗРЕЧЕНИЯ"), 7.5, bold=True, color=ACCENT)
    text = left.add_paragraph()
    spacing(text, after=4, line=1.1)
    set_font(text.add_run(data.get("day_in_three", "")), 9.5)

    if weather:
        right = table.rows[0].cells[1]
        head = right.paragraphs[0]
        spacing(head, before=2, after=1)
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(head.add_run("ВРЕМЕТО ДНЕС · " + weather["city"].upper()), 7.5,
                 bold=True, color=ACCENT)

        temps = right.add_paragraph()
        spacing(temps, after=0)
        temps.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(temps.add_run("{}°".format(weather["max"])), 22, bold=True)
        set_font(temps.add_run(" / {}°".format(weather["min"])), 14, color=MUTED)

        for line, size, bold in (
                (weather["text"], 9.5, True),
                ("Сутрин {}° · Обед {}° · Вечер {}°".format(
                    weather["morning"], weather["noon"], weather["evening"]), 8, False),
                ("Дъжд {}% · Вятър до {} км/ч".format(weather["rain"], weather["wind"]), 8, False),
                (weather_tip(weather), 8, True)):
            p = right.add_paragraph()
            spacing(p, after=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(line), size, bold=bold)


def build_docx(data, items, today, weather=None):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(1.0))
    section.top_margin = section.bottom_margin = Cm(0.9)

    # Глава на вестника
    kicker = doc.add_paragraph()
    spacing(kicker, after=2)
    add_border(kicker, "bottom", size=4, color="CFC6B4")
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(kicker.add_run("{}  ·  {} истории  ·  съставен в {}".format(
        bg_date(today), len(data["stories"]), datetime.now().strftime("%H:%M"))),
        7.5, color=MUTED)

    name = doc.add_paragraph()
    spacing(name, after=2)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_border(name, "bottom", size=12, style="double", space=2)
    set_font(name.add_run("Сутрешен вестник"), 26, bold=True)

    add_day_box(doc, data, weather)

    # Новините в две колони
    body = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body, 2)
    # Празният параграф, който носи прехода към колоните, да не заема място.
    breaker = doc.paragraphs[-1]
    spacing(breaker, after=3, line=Pt(1))

    stories = data["stories"]
    top = [s for s in stories if s.get("top")]
    if top:
        add_heading(doc, "Главното днес", ACCENT)
        for s in top:
            add_story(doc, s, items, big=True)
    for sec_name in SECTIONS:
        group = [s for s in stories if not s.get("top") and s.get("section") == sec_name]
        if group:
            add_heading(doc, sec_name)
            for s in group:
                add_story(doc, s, items)
    return doc


# ---------- main ----------

DESKTOP_COPY = Path.home() / "Desktop" / "Сутрешен вестник.docx"


def notify(message):
    """Малко известие в ъгъла на екрана на Mac."""
    script = 'display notification "{}" with title "Сутрешен вестник"'.format(
        message.replace('"', "'"))
    subprocess.run(["osascript", "-e", script], capture_output=True)


def wait_for_internet(max_wait=120):
    """След събуждане Wi-Fi се свързва за няколко секунди, затова изчакваме."""
    deadline = time.time() + max_wait
    while True:
        try:
            urllib.request.urlopen("https://api.open-meteo.com", timeout=5)
            return
        except urllib.error.HTTPError:
            return  # сървърът отговори, значи има интернет
        except Exception:
            if time.time() > deadline:
                log.warning("Няма интернет след %d секунди, опитваме все пак", max_wait)
                return
            time.sleep(5)


def main():
    auto = "--auto" in sys.argv
    today = datetime.now()
    out = NEWSPAPERS / "{}.docx".format(today.strftime("%Y-%m-%d"))
    if auto and "--force" not in sys.argv and (out.exists() or today.hour < 5):
        return  # днешният брой вече е готов или е твърде рано

    setup_logging()
    log.info("=== Нов брой%s ===", " (автоматично)" if auto else "")
    try:
        NEWSPAPERS.mkdir(exist_ok=True)
        if auto:
            wait_for_internet()
        cache = NEWSPAPERS / "{}.json".format(today.strftime("%Y-%m-%d"))
        config = load_config()
        place = config.get("weather")
        weather = fetch_weather(place) if place else None

        if "--redo" in sys.argv and cache.exists():
            # Само прередактира документа от вече обобщените новини, без нов Claude.
            saved = json.loads(cache.read_text(encoding="utf-8"))
            data, items = saved["data"], saved["items"]
        else:
            items = fetch_all(config)
            if not items:
                raise RuntimeError("Нито един източник не върна новини")
            items = dedupe(items)
            log.info("Общо %d новини след махане на дубликати", len(items))
            data = summarize(items, config["settings"])
            cache.write_text(json.dumps({"data": data, "items": items}, ensure_ascii=False),
                             encoding="utf-8")

        build_docx(data, items, today, weather).save(out)
        shutil.copyfile(out, BASE / "latest.docx")
        shutil.copyfile(out, DESKTOP_COPY)
        log.info("Готово: %s (%d истории)", out, len(data["stories"]))
        if auto:
            notify("Готов е! {} истории — файлът е на Desktop.".format(len(data["stories"])))
    except Exception as e:
        log.exception("Вестникът не беше създаден")
        if auto:
            notify("Грешка: {}. Виж logs/vestnik.log".format(str(e)[:80]))
        sys.exit(1)


if __name__ == "__main__":
    main()
